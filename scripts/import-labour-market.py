#!/usr/bin/env python3
"""
Imports Jobs and Skills Australia labour market workbooks into Bcome's /data folder.

    pip install openpyxl
    python3 scripts/import-labour-market.py \
        --projections  "employment_projections_-_may_2025_to_may_2035.xlsx" \
        --tnv          "total_new_vacancies_may_2026.xlsx" \
        --ivi          "internet_vacancies_anzsco_skill_level_states_and_territories_-_august_2026.xlsx" \
        --industry     "national_industry_trend_-_february_2026.xlsx" \
        --osl          "2025_Key_Findings_Report_downloadable_Charts_and_Tables_0.xlsx" \
        --genai        "jsa_gen_ai_interactive_table_data_pack_20250903.xlsx"

Any of the options can be left out to refresh only some files. Writes:
  data/outlook/jsa-2025.json      projections for every ANZSCO unit group, major group, skill level and industry
  data/labour/vacancies.json      Total New Vacancies (quarterly, ANZSCO 1 and 2 digit) and the
                                  Internet Vacancy Index by skill level and state (monthly, trend)
  data/labour/industries.json     employment by ANZSIC division, subdivision and group (trend)
  data/labour/shortage.json       Occupation Shortage List 2025 key findings
  data/labour/genai.json          Gen AI Capacity Study exposure scores by ANZSCO unit group and industry

Then run `npm run data:validate` and `npm run db:seed-sql`.
"""
import argparse
import datetime as dt
import json
import os
import re
import sys

try:
    import openpyxl
except ImportError:
    sys.exit("Install openpyxl first: pip install openpyxl")

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
TODAY = dt.date.today().isoformat()


def rows(path, sheet):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    return [list(r) for r in wb[sheet].iter_rows(values_only=True)]


def sheet_like(path, pattern):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for name in wb.sheetnames:
        if re.search(pattern, name, re.I):
            return name
    sys.exit(f"{os.path.basename(path)}: no sheet matching {pattern} (found {wb.sheetnames})")


def r1(x):
    return None if x is None else round(float(x), 1)


def pct(x):
    """JSA stores percentages as fractions in the projections workbook."""
    return None if x is None or isinstance(x, str) else round(float(x) * 100, 1)


def num(x):
    return float(x) if isinstance(x, (int, float)) else None


def write(rel, data):
    path = os.path.join(ROOT, rel)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w") as f:
        json.dump(data, f, indent=1, ensure_ascii=False)
        f.write("\n")
    print(f"wrote {rel} ({os.path.getsize(path) // 1024} KB)")


def ym(d):
    return d.strftime("%Y-%m") if isinstance(d, (dt.date, dt.datetime)) else str(d)[:7]


# ---------------------------------------------------------------------------
# Employment projections (May 2025 to May 2035)
# ---------------------------------------------------------------------------

def skill_levels(v):
    if v in (None, "-"):
        return []
    return sorted({int(s) for s in re.findall(r"[1-5]", str(v))})


def import_projections(path):
    out_path = os.path.join(ROOT, "data/outlook/jsa-2025.json")
    jsa = json.load(open(out_path))

    t6 = rows(path, sheet_like(path, r"Table_6"))
    occupations = {}
    for r in t6:
        if r[0] == 4 and r[1] == "N" and re.fullmatch(r"\d{4}", str(r[2])):
            occupations[str(r[2])] = {
                "name": r[3].strip(),
                "skillLevels": skill_levels(r[4]),
                "employed2025k": r1(r[5]),
                "employed2030k": r1(r[6]),
                "employed2035k": r1(r[7]),
                "growth5Pct": pct(r[9]),
                "growth10Pct": pct(r[11]),
            }
    # Sub-major (2-digit) and minor (3-digit) groups, for vacancy estimates.
    groups = {}
    for r in t6:
        if r[0] in (2, 3) and r[1] == "N" and isinstance(r[5], (int, float)):
            groups[str(r[2])] = {"name": r[3].strip(), "employed2025k": r1(r[5]), "growth10Pct": pct(r[11])}

    t2 = rows(path, sheet_like(path, r"Table_2"))
    majors = {}
    for r in t2:
        if r[0] == 1 and isinstance(r[3], (int, float)):
            majors[str(r[1])] = {"name": r[2].strip(), "employed2025k": r1(r[3]), "growth5Pct": pct(r[10]), "growth10Pct": pct(r[12])}
        if r[2] == "Total":
            jsa["total"] = {"employed2025k": r1(r[3]), "employed2035k": r1(r[5]), "growth5Pct": pct(r[10]), "growth10Pct": pct(r[12]), "growth10k": r1(r[11])}

    t3 = rows(path, sheet_like(path, r"Table_3"))
    labels = {"1": "Bachelor degree or higher", "2": "Advanced diploma or diploma", "3": "Certificate IV or III (with experience)", "4": "Certificate II or III", "5": "Certificate I or secondary education"}
    levels = {}
    for r in t3:
        m = re.fullmatch(r"Skill level (\d)", str(r[0] or ""))
        if m:
            levels[m.group(1)] = {"label": labels[m.group(1)], "employed2025k": r1(r[1]), "growth5Pct": pct(r[8]), "growth10Pct": pct(r[10])}

    t1 = rows(path, sheet_like(path, r"Table_1"))
    # Keep the division names already used by occupations (JSA spells a couple differently).
    canonical = {"Healthcare and Social Assistance": "Health Care and Social Assistance", "Retail trade": "Retail Trade"}
    industries, divisions = {}, {}
    for r in t1:
        if r[0] == 1 and isinstance(r[3], (int, float)):
            name = canonical.get(r[2].strip(), r[2].strip())
            industries[name] = round(pct(r[12]))
            divisions[str(r[1])] = {"name": name, "employed2025k": r1(r[3]), "employed2035k": r1(r[5]), "growth5Pct": pct(r[10]), "growth10Pct": pct(r[12])}

    jsa["majorGroups"] = majors
    jsa["occupationGroups"] = groups
    jsa["occupations"] = occupations
    jsa["skillLevels"] = levels
    jsa["industries"] = industries
    jsa["divisions"] = divisions
    jsa["source"]["retrieved"] = TODAY
    jsa["source"]["coverage"] = f"All {len(occupations)} ANZSCO unit groups (Table 6), major groups, skill levels and industry divisions, imported from the JSA workbook with scripts/import-labour-market.py."
    jsa["source"]["classification"] = "ANZSCO version 1.3 and ANZSIC 2006 (Revision 2.0), as published by JSA"
    jsa["skillLevelNote"] = "Skill levels are set by ANZSCO for each occupation. Where a unit group contains occupations at different skill levels, JSA lists each of them."
    write("data/outlook/jsa-2025.json", jsa)


# ---------------------------------------------------------------------------
# Total New Vacancies (quarterly) and the Internet Vacancy Index (monthly)
# ---------------------------------------------------------------------------

STATE_CODES = {"1": "NSW", "2": "VIC", "3": "QLD", "4": "SA", "5": "WA", "6": "TAS", "7": "NT", "8": "ACT"}


def import_vacancies(tnv_path, ivi_path):
    out_path = os.path.join(ROOT, "data/labour/vacancies.json")
    out = json.load(open(out_path)) if os.path.exists(out_path) else {"source": {}}

    if tnv_path:
        d2 = rows(tnv_path, "Data_2")
        header = next(i for i, r in enumerate(d2) if r[:2] == ["date", "occupation_code"])
        series, titles, quarters = {}, {}, set()
        for r in d2[header + 1:]:
            if r[0] is None or r[4] is None:
                continue
            q = ym(r[0])
            quarters.add(q)
            code = str(r[1])
            titles[code] = (r[2], r[3])
            series.setdefault(code, {})[q] = int(r[4])
        quarters = sorted(quarters)
        out["tnv"] = {
            "quarters": quarters,
            "latest": quarters[-1],
            "occupations": {
                code: {"level": lvl, "title": title, "values": [series[code].get(q) for q in quarters]}
                for code, (lvl, title) in sorted(titles.items(), key=lambda kv: (len(kv[0]), kv[0]))
            },
        }
        d1 = rows(tnv_path, "Data_1")
        header = next(i for i, r in enumerate(d1) if r[:2] == ["region_level", "region_code"])
        regions, rq = {}, set()
        for r in d1[header + 1:]:
            if r[0] in ("national", "state") and r[4] is not None:
                code = "AUST" if r[0] == "national" else STATE_CODES.get(str(r[1]), str(r[1]))
                q = ym(r[3])
                rq.add(q)
                regions.setdefault(code, {})[q] = int(r[4])
        rq = sorted(rq)
        out["tnv"]["regionQuarters"] = rq
        out["tnv"]["regions"] = {code: [v.get(q) for q in rq] for code, v in regions.items()}
        out["source"]["tnv"] = {
            "title": "Jobs and Skills Australia, Total New Vacancies (TNV), May 2026",
            "url": "https://www.jobsandskills.gov.au/data/total-new-vacancies",
            "note": "New vacancies created in each quarter, including jobs not advertised online (REOS employer survey combined with the Internet Vacancy Index). JSA advises comparing years rather than single quarters.",
            "retrieved": TODAY,
        }

    if ivi_path:
        t = rows(ivi_path, "Trend")
        header = next(i for i, r in enumerate(t) if r[:4] == ["Level", "Title", "State", "Skill_level"])
        months = [ym(c) for c in t[header][4:] if c is not None]
        ivi = {}
        for r in t[header + 1:]:
            if r[2] is None:
                continue
            state = "AUST" if r[2] == "AUST" else r[2]
            vals = [None if v is None else int(round(v)) for v in r[4 : 4 + len(months)]]
            ivi.setdefault(state, {})[str(int(r[3]))] = vals
        out["ivi"] = {"months": months, "latest": months[-1], "series": ivi}
        out["source"]["ivi"] = {
            "title": "Jobs and Skills Australia, Internet Vacancy Index, ANZSCO skill level, states and territories, August 2026 (trend)",
            "url": "https://www.jobsandskills.gov.au/data/internet-vacancy-index",
            "note": "Online job advertisements newly lodged each month on SEEK, Workforce Australia and CareerOne, by the skill level of the occupation. Trend series.",
            "retrieved": TODAY,
        }
    write("data/labour/vacancies.json", out)


# ---------------------------------------------------------------------------
# Employment by industry (ANZSIC), trend
# ---------------------------------------------------------------------------

def import_industries(path):
    t1 = rows(path, "Table_1")
    header = next(i for i, r in enumerate(t1) if r[0] == "ANZSIC Level")
    summary = {}
    for r in t1[header + 1:]:
        if r[0] not in (1, 2, 3) or r[2] is None:
            continue
        summary[str(r[2])] = {
            "level": r[0],
            "nfd": r[1] == "Y",
            "title": r[3].strip(),
            "division": r[5],
            "employedK": r1(num(r[6])),
            "change1yPct": r1(num(r[10])),
            "change5yPct": r1(num(r[12])),
            "change10yPct": r1(num(r[14])),
            "sharePct": r1(num(r[15])),
            "femalePct": r1(num(r[17])),
            "partTimePct": r1(num(r[18])),
            "medianAge": num(r[21]) if len(r) > 21 else None,
        }
    t2 = rows(path, "Table_2")
    h2 = next(i for i, r in enumerate(t2) if r[0] == "ANZSIC Level")
    dates = [ym(c) for c in t2[h2][5:] if c is not None]
    # Annual points (each February) from 2006, so the history chart stays light.
    keep = [i for i, d in enumerate(dates) if d.endswith("-02") and d >= "2006-02"]
    years = [dates[i][:4] for i in keep]
    for r in t2[h2 + 1:]:
        code = str(r[2]) if r[2] is not None else None
        if code in summary:
            vals = r[5 : 5 + len(dates)]
            summary[code]["series"] = [r1(num(vals[i])) for i in keep]
    out = {
        "source": {
            "title": "Jobs and Skills Australia, Labour Force Trending: National Industry Trend, February 2026",
            "url": "https://www.jobsandskills.gov.au/data/labour-market-insights/industries",
            "note": "Employment ('000) by ANZSIC 2006 division, subdivision and group, from the ABS Labour Force Survey, trended by JSA. Gender, part-time and median age are 2025 four-quarter averages.",
            "reference": ym(dates[-1]),
            "retrieved": TODAY,
        },
        "years": years,
        "industries": summary,
    }
    write("data/labour/industries.json", out)


# ---------------------------------------------------------------------------
# Occupation Shortage List 2025 key findings
# ---------------------------------------------------------------------------

def table(path, sheet, first_col):
    t = rows(path, sheet)
    h = next(i for i, r in enumerate(t) if r[0] == first_col)
    head = [c for c in t[h] if c is not None]
    body = []
    for r in t[h + 1:]:
        if r[0] is None:
            break
        body.append(r[: len(head)])
    return head, body


def import_osl(path):
    def by_year(sheet, first):
        head, body = table(path, sheet, first)
        years = [str(y) for y in head[1:]]
        return years, {re.sub(r"\s+", " ", str(r[0])).replace("&", "and").strip(): [r1(v * 100) for v in r[1:]] for r in body}

    years, major = by_year("Figure 2", "Major Group")
    _, skill = by_year("Figure 4", "Skill level")
    head, apps = table(path, "Figure 7", "Applicants per vacancy (no.)")
    out = {
        "source": {
            "title": "Jobs and Skills Australia, 2025 Occupation Shortage List Key Findings Report (charts and tables)",
            "url": "https://www.jobsandskills.gov.au/data/occupation-shortages-analysis/occupation-shortage-list",
            "note": "Share of occupations rated in shortage, OSCA basis. The 2025 OSL assessment covers April 2024 to June 2025.",
            "retrieved": TODAY,
        },
        "years": years,
        "byMajorGroup": major,
        "bySkillLevel": skill,
        "applicantsPerVacancy": {"years": [str(y) for y in head[1:]], **{str(r[0]): [round(float(v), 1) for v in r[1:]] for r in apps}},
    }
    write("data/labour/shortage.json", out)


# ---------------------------------------------------------------------------
# Gen AI Capacity Study ("Our Gen AI Transition"), interactive tables data pack
# ---------------------------------------------------------------------------

def opt(x, digits=3):
    return round(float(x), digits) if isinstance(x, (int, float)) else None


def tercile_cuts(values):
    """The scores a third and two thirds of the way up, so the bands split occupations evenly."""
    v = sorted(values)
    return [round(v[len(v) // 3], 3), round(v[(2 * len(v)) // 3], 3)]


def import_genai(path):
    t = rows(path, sheet_like(path, r"^occupation"))
    h = next(i for i, r in enumerate(t) if r[0] and str(r[0]).strip().lower() == "anzsco unit code")
    head = [re.sub(r"\s+", " ", str(c or "")).strip().lower() for c in t[h]]

    def col(pattern):
        for i, c in enumerate(head):
            if re.search(pattern, c):
                return i
        sys.exit(f"{os.path.basename(path)}: no column matching {pattern} in {head}")

    c = {
        "title": col(r"unit title"),
        "group": col(r"matrix group"),
        "aug": col(r"^augmentation exposure"),
        "augSd": col(r"^augmentation standard"),
        "auto": col(r"^automation exposure"),
        "autoSd": col(r"^automation standard"),
        "skillChange": col(r"rate of skill change"),
        "mobility": col(r"historical occupation mobility"),
        "highFit": col(r"high-fit transition"),
        "hybrid": col(r"hybridisation"),
        "special": col(r"specialisation"),
        "entry": col(r"entry level"),
    }
    occ = {}
    for r in t[h + 1:]:
        if r[0] is None or not re.fullmatch(r"\d{4}", str(r[0]).strip()):
            continue
        if not isinstance(r[c["auto"]], (int, float)) or not isinstance(r[c["aug"]], (int, float)):
            continue
        entry = opt(r[c["entry"]])
        occ[str(r[0]).strip()] = {
            "title": str(r[c["title"]]).strip(),
            "group": str(r[c["group"]] or "").strip() or None,
            "automation": opt(r[c["auto"]]),
            "automationSd": opt(r[c["autoSd"]]),
            "augmentation": opt(r[c["aug"]]),
            "augmentationSd": opt(r[c["augSd"]]),
            "skillChange": opt(r[c["skillChange"]], 1),
            "mobility": opt(r[c["mobility"]]),
            "highFitTransitionPct": None if opt(r[c["highFit"]]) is None else round(opt(r[c["highFit"]]) * 100, 1),
            "hybridisation": opt(r[c["hybrid"]]),
            "specialisation": opt(r[c["special"]]),
            "entryLevelAdsPct": None if entry is None else round(entry * 100, 1),
        }
    if len(occ) < 300:
        sys.exit(f"{os.path.basename(path)}: only {len(occ)} occupations found; has the layout changed?")

    ti = rows(path, sheet_like(path, r"^industry"))
    hi = next(i for i, r in enumerate(ti) if r[0] and "anzsic division" in str(r[0]).lower())
    industries = {}
    for r in ti[hi + 1:]:
        if r[0] is None or not isinstance(r[1], (int, float)):
            continue
        industries[str(r[0]).strip()] = {
            "augmentation": opt(r[1]),
            "automation": opt(r[2]),
            "aiSkillsWorkerProfilesPct": None if opt(r[3], 5) is None else round(r[3] * 100, 2),
            "aiSkillsJobAdsPct": None if opt(r[4], 5) is None else round(r[4] * 100, 2),
            "entryLevelAdsPct": None if opt(r[5]) is None else round(r[5] * 100, 1),
        }

    auto = [o["automation"] for o in occ.values()]
    aug = [o["augmentation"] for o in occ.values()]
    m = re.search(r"(20\d{2})(\d{2})(\d{2})", os.path.basename(path))
    out = {
        "source": {
            "title": "Jobs and Skills Australia, Our Gen AI Transition (Generative AI Capacity Study), interactive tables data pack",
            "url": "https://www.jobsandskills.gov.au/studies/generative-artificial-intelligence-capacity-study",
            "file": os.path.basename(path),
            "released": f"{m.group(1)}-{m.group(2)}-{m.group(3)}" if m else None,
            "note": "Task-based scores from 0 to 1 for how much of an occupation's work Gen AI could automate (do) or augment (help with), assuming it is fully adopted. JSA does not set high, medium or low thresholds; Bcome splits occupations into thirds by score. JSA's modelling looks at adoption to 2035 and 2050.",
            "retrieved": TODAY,
        },
        "bands": {"automation": tercile_cuts(auto), "augmentation": tercile_cuts(aug)},
        "occupations": dict(sorted(occ.items())),
        "industries": industries,
    }
    write("data/labour/genai.json", out)


if __name__ == "__main__":
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--projections")
    p.add_argument("--tnv")
    p.add_argument("--ivi")
    p.add_argument("--industry")
    p.add_argument("--osl")
    p.add_argument("--genai")
    a = p.parse_args()
    if not any(vars(a).values()):
        p.print_help()
        sys.exit(1)
    if a.projections:
        import_projections(a.projections)
    if a.tnv or a.ivi:
        import_vacancies(a.tnv, a.ivi)
    if a.industry:
        import_industries(a.industry)
    if a.osl:
        import_osl(a.osl)
    if a.genai:
        import_genai(a.genai)
